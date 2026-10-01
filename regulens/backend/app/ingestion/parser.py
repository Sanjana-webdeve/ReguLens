"""Text extraction for PDF / DOCX / TXT. Returns pages -> paragraphs and raises ParseError with a user-friendly message."""
from dataclasses import dataclass


class ParseError(Exception):
    pass


@dataclass
class Page:
    number: int
    paragraphs: list[str]


def parse_file(path: str, ext: str) -> list[Page]:
    try:
        if ext == "pdf":
            pages = _parse_pdf(path)
        elif ext == "docx":
            pages = _parse_docx(path)
        elif ext == "txt":
            pages = _parse_txt(path)
        else:
            raise ParseError(f"Unsupported file format '.{ext}'. Upload PDF, DOCX or TXT.")
    except ParseError:
        raise
    except Exception as e:
        raise ParseError("The uploaded file could not be parsed. It may be corrupted or password-protected.") from e
    if not any(p.paragraphs for p in pages):
        raise ParseError("No extractable text was found. Scanned/image-only documents need OCR, which is not enabled.")
    return pages


def _parse_pdf(path: str) -> list[Page]:
    import fitz  # PyMuPDF

    pages = []
    with fitz.open(path) as doc:
        if doc.needs_pass:
            raise ParseError("The PDF is password-protected.")
        for i, page in enumerate(doc, start=1):
            blocks = page.get_text("blocks", sort=True)
            paras = [b[4].strip() for b in blocks if b[6] == 0 and b[4].strip()]
            pages.append(Page(i, paras))
    return pages


def _parse_docx(path: str) -> list[Page]:
    import docx

    d = docx.Document(path)
    paras = []
    for p in d.paragraphs:
        t = p.text.strip()
        if t:
            paras.append(t)
    # DOCX has no fixed pages: emulate ~3000 characters per "page"
    pages, cur, size = [], [], 0
    for t in paras:
        cur.append(t)
        size += len(t)
        if size >= 3000:
            pages.append(Page(len(pages) + 1, cur))
            cur, size = [], 0
    if cur:
        pages.append(Page(len(pages) + 1, cur))
    return pages


def _parse_txt(path: str) -> list[Page]:
    with open(path, "rb") as f:
        raw = f.read()
    text = raw.decode("utf-8", errors="replace")
    paras = [p.strip() for p in text.replace("\r", "").split("\n\n") if p.strip()]
    pages, cur, size = [], [], 0
    for t in paras:
        cur.append(t)
        size += len(t)
        if size >= 3000:
            pages.append(Page(len(pages) + 1, cur))
            cur, size = [], 0
    if cur:
        pages.append(Page(len(pages) + 1, cur))
    return pages
